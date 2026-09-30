import csv,re
from pathlib import Path
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]; X=R/'new data.xlsx'; D=R/'h301194_school.sql'
def norm(v):
    if v is None:return ''
    s=str(v).strip().replace('ي','ی').replace('ى','ی').replace('ك','ک').replace('\u200c',' ')
    s=re.sub(r'[\u064B-\u065F\u0670]','',s); s=re.sub(r'[\s\-_/]+',' ',s)
    return s.strip().casefold()
def unq(v):
    v=v.strip()
    if v.upper()=='NULL':return None
    if len(v)>=2 and v[0]=="'" and v[-1]=="'":return v[1:-1].replace("''","'")
    return v
def tuples(block):
    out=[]; cur=[]; t=''; q=False; esc=False; dep=0
    for ch in block:
        if q:
            t+=ch
            if esc:esc=False
            elif ch=='\\':esc=True
            elif ch=="'":q=False
        elif ch=="'":q=True;t+=ch
        elif ch=='(':dep+=1;cur=[];t=''
        elif ch==',' and dep==1:cur.append(t.strip())
        elif ch==')' and dep==1:cur.append(t.strip());out.append(cur);dep=0;t=''
        elif dep:t+=ch
    return out
text=D.read_text(encoding='utf-8',errors='replace')
m=re.search(r'INSERT INTO [`]?schools[`]?\s*\((.*?)\)\s*VALUES\s*(.*?);',text,re.S|re.I)
cols=[x.strip().strip('` ') for x in m.group(1).split(',')]; db=[]
for row in tuples(m.group(2)):
    if len(row)==len(cols):db.append({c:unq(row[i]) for i,c in enumerate(cols)})
md=re.search(r'INSERT INTO [`]?districts[`]?\s*\((.*?)\)\s*VALUES\s*(.*?);',text,re.S|re.I)
districts={}
if md:
    dc=[x.strip().strip('` ') for x in md.group(1).split(',')]
    for row in tuples(md.group(2)):
        if len(row)==len(dc):
            z={c:unq(row[i]) for i,c in enumerate(dc)}; districts[norm(z.get('id'))]=z.get('title')
for d in db:d['_district']=districts.get(norm(d.get('district_id')),'');d['_key']=(norm(d.get('name')),norm(d.get('_district')),norm(d.get('level')))
wb=load_workbook(X,data_only=True,read_only=True); excel=[]
def col(h,names):
    ns={norm(x) for x in names}
    for i,x in enumerate(h):
        if norm(x) in ns:return i
for ws in wb.worksheets:
    rr=[list(x) for x in ws.iter_rows(values_only=True) if any(v not in (None,'') for v in x)]
    if not rr:continue
    h=rr[0]; C={
      'code':col(h,['کد مدرسه','کد','school code']),
      'name':col(h,['نام مدرسه','نام','school name']),
      'district':col(h,['ناحیه','منطقه','district']),
      'level':col(h,['مقطع','مقطع تحصیلی','level']),
      'company':col(h,['شرکت','شرکت سرویس','company']),
      'students':col(h,['تعداد دانش آموز','تعداد دانش‌آموز','دانش آموز','student count'])}
    for row in rr[1:]:
        g=lambda k: row[C[k]] if C[k] is not None and C[k]<len(row) else None
        if g('name') and g('district'):excel.append({'sheet':ws.title,'code':g('code'),'name':g('name'),'district':g('district'),'level':g('level'),'company':g('company'),'students':g('students')})
bycode={norm(d.get('code')):d for d in db}; bykey={}
for d in db:bykey.setdefault(d['_key'],[]).append(d)
matched=[];amb=[];missing=[];seen=set()
for e in excel:
    cs=[]; c=norm(e['code'])
    if c and c in bycode:
        d=bycode[c]
        if d['_key']==(norm(e['name']),norm(e['district']),norm(e['level'])):cs=[d]
    if not cs:cs=bykey.get((norm(e['name']),norm(e['district']),norm(e['level'])),[])
    if len(cs)==1:matched.append((e,cs[0]));seen.add(str(cs[0]['id']))
    elif len(cs)>1:amb.append((e,cs))
    else:missing.append(e)
def q(v):
    if v is None or str(v).strip()=='':return 'NULL'
    return "'"+str(v).replace('\\','\\\\').replace("'","''")+"'"
def n(v):
    if v is None or str(v).strip()=='':return 'NULL'
    try:return str(int(float(str(v).replace(',','').strip())))
    except:return 'NULL'
L=['START TRANSACTION;']
for e,d in matched:
    sets=['is_active=1','student_count='+n(e['students'])]
    if e['company'] not in (None,''):sets.append('company_id=(SELECT id FROM companies WHERE title='+q(e['company'])+' ORDER BY id LIMIT 1)')
    L.append('UPDATE schools SET '+', '.join(sets)+' WHERE id='+str(d['id'])+';')
for d in db:
    if str(d['id']) not in seen:L.append('UPDATE schools SET is_active=0 WHERE id='+str(d['id'])+';')
L+=['COMMIT;']
(R/'generated_school_sync.sql').write_text('\n'.join(L)+'\n',encoding='utf-8')
with (R/'school_sync_audit.csv').open('w',newline='',encoding='utf-8-sig') as f:
    w=csv.writer(f);w.writerow(['status','sheet','name','district','level','code','company','students','db_id','note'])
    for e,d in matched:w.writerow(['matched',e['sheet'],e['name'],e['district'],e['level'],e['code'],e['company'],e['students'],d['id'],''])
    for e,c in amb:w.writerow(['ambiguous',e['sheet'],e['name'],e['district'],e['level'],e['code'],e['company'],e['students'],'','multiple candidates'])
    for e in missing:w.writerow(['missing_in_db',e['sheet'],e['name'],e['district'],e['level'],e['code'],e['company'],e['students'],'','no DB candidate'])
(R/'school_sync_summary.txt').write_text('excel='+str(len(excel))+' matched='+str(len(matched))+' ambiguous='+str(len(amb))+' missing_in_db='+str(len(missing))+' db='+str(len(db))+' deactivate='+str(len(db)-len(seen))+'\n',encoding='utf-8')